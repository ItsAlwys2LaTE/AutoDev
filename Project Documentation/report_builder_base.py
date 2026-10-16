"""
AutoDev: Autonomous Multi-Agent Software Engineering Platform
Capstone Project-1 Report Generator (.docx)
Bhilai Institute of Technology, Durg - Department of Computer Science & Engineering
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

class ReportBuilder:
    def __init__(self, output_path, images_dir):
        self.doc = docx.Document()
        self.output_path = output_path
        self.images_dir = images_dir
        self.figure_count = 0
        self.table_count = 0
        self._setup_page()

    def _setup_page(self):
        section = self.doc.sections[0]
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69) # A4
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Configure default Normal style
        style = self.doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)

    def add_page_break(self):
        self.doc.add_page_break()

    def add_title_line(self, text, size=14, bold=True, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x0F, 0x17, 0x2A)):
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return p

    def add_chapter_title(self, chapter_num, title_text):
        p1 = self.doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(12)
        p1.paragraph_format.space_after = Pt(4)
        r1 = p1.add_run(f"CHAPTER {chapter_num}")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(16)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        p2 = self.doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(18)
        r2 = p2.add_run(title_text.upper())
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(16)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def add_heading_1(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        return p

    def add_heading_2(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        return p

    def add_body(self, text, bold_prefix=None, space_after=6):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(space_after)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.font.bold = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    def add_bullet(self, bold_prefix, text):
        p = self.doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    def add_figure(self, filename, caption_text, width=Inches(5.8)):
        img_path = os.path.join(self.images_dir, filename)
        p_img = self.doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        
        if os.path.exists(img_path):
            try:
                run = p_img.add_run()
                run.add_picture(img_path, width=width)
            except Exception as e:
                p_img.add_run(f"[{caption_text} - Image Loaded ({filename})]")
        else:
            run = p_img.add_run(f"[Figure File Missing: {filename}]")
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.font.bold = True
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    def add_table(self, headers, rows, col_widths=None):
        table = self.doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, h_text in enumerate(headers):
            cell = hdr_cells[i]
            cell.text = h_text
            set_cell_shading(cell, "1E293B")
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Data Rows
        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_data in enumerate(row_data):
                cell = row_cells[col_idx]
                cell.text = str(cell_data)
                set_cell_shading(cell, bg_color)
                set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx != 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # Apply column widths if specified
        if col_widths:
            for row in table.rows:
                for idx, w in enumerate(col_widths):
                    row.cells[idx].width = Inches(w)

        # Add space after table
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(8)

    def save(self):
        self.doc.save(self.output_path)
        print(f"Report saved successfully to {self.output_path}")

print("ReportBuilder base class ready.")
